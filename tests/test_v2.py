"""第二版功能測試：自動儲存、排序、QR Code、批次匯入、發布歷史、瀏覽統計、帳號期限、公告、備份、縮圖、SEO。

執行：python -m unittest discover -s tests -t . -v
"""
import io
import json
import re
import zipfile

from tests.test_app import Base, png_bytes


class AutosaveTests(Base):
    def test_profile_autosave_saves_text_and_returns_time(self):
        self.login("t001")
        r = self.post("/dashboard/autosave/profile", {"county": "新北市", "school_name": "自動儲存學校",
                                                      "teacher_name": "林怡君", "slogan": "自動儲存 slogan"})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.get_json()["ok"])
        self.assertRegex(r.get_json()["saved_at"], r"^\d\d:\d\d$")
        self.assertEqual(self.db("SELECT slogan FROM teacher_profiles WHERE teacher_code='001'")[0][0], "自動儲存 slogan")
        # 自動儲存只寫草稿，不影響公開頁
        self.assertNotIn("自動儲存 slogan", self.app.test_client().get("/teacher/001").get_data(as_text=True))

    def test_autosave_validation_and_csrf(self):
        self.login("t001")
        r = self.post("/dashboard/autosave/profile", {"county": "火星", "teacher_name": "x"})
        self.assertEqual(r.status_code, 400)
        self.assertFalse(r.get_json()["ok"])
        r = self.client.post("/dashboard/autosave/profile", data={"slogan": "no token"})
        self.assertEqual(r.status_code, 400)

    def test_course_autosave_ownership(self):
        self.login("t001")
        own, other = self.course_of("t001"), self.course_of("t002")
        r = self.post(f"/dashboard/autosave/course/{own}", {"title": "自動儲存課程", "ai_usage": "新的 AI 運用"})
        self.assertTrue(r.get_json()["ok"])
        r = self.post(f"/dashboard/autosave/course/{other}", {"title": "HACK"})
        self.assertEqual(r.status_code, 404)
        self.assertNotEqual(self.db("SELECT title FROM courses WHERE id=%s", (other,))[0][0], "HACK")


class ReorderTests(Base):
    def test_reorder_own_items(self):
        self.login("t001")
        ids = [r[0] for r in self.db("SELECT s.id FROM specialties s JOIN teacher_profiles p ON p.id=s.profile_id"
                                     " WHERE p.teacher_code='001' ORDER BY s.sort_order")]
        new = list(reversed(ids))
        r = self.post("/dashboard/reorder/specialties", {"ids": ",".join(map(str, new))})
        self.assertTrue(r.get_json()["ok"])
        got = [r[0] for r in self.db("SELECT s.id FROM specialties s JOIN teacher_profiles p ON p.id=s.profile_id"
                                     " WHERE p.teacher_code='001' ORDER BY s.sort_order")]
        self.assertEqual(got, new)

    def test_reorder_rejects_foreign_or_partial_ids(self):
        self.login("t001")
        own = [r[0] for r in self.db("SELECT ph.id FROM teacher_photos ph JOIN teacher_profiles p ON p.id=ph.profile_id"
                                     " WHERE p.teacher_code='001'")]
        foreign = self.db("SELECT ph.id FROM teacher_photos ph JOIN teacher_profiles p ON p.id=ph.profile_id"
                          " WHERE p.teacher_code='002'")[0][0]
        r = self.post("/dashboard/reorder/photos", {"ids": ",".join(map(str, own[:-1] + [foreign]))})
        self.assertEqual(r.status_code, 409)
        r = self.post("/dashboard/reorder/photos", {"ids": ",".join(map(str, own[:-1]))})
        self.assertEqual(r.status_code, 409)
        self.assertEqual(self.post("/dashboard/reorder/users", {"ids": "1"}).status_code, 404)

    def test_reorder_courses(self):
        self.login("t003")
        for t in ("A 課", "B 課"):
            self.post("/dashboard/courses/new", {"title": t})
        ids = [r[0] for r in self.db("SELECT c.id FROM courses c JOIN teacher_profiles p ON p.id=c.profile_id"
                                     " WHERE p.teacher_code='003' ORDER BY c.sort_order")]
        self.post("/dashboard/reorder/courses", {"ids": f"{ids[1]},{ids[0]}"})
        self.assertEqual(self.db("SELECT title FROM courses c JOIN teacher_profiles p ON p.id=c.profile_id"
                                 " WHERE p.teacher_code='003' ORDER BY c.sort_order")[0][0], "B 課")


class QrTests(Base):
    def test_qr_decodes_to_public_url(self):
        try:
            import cv2
            import numpy as np
            from PIL import Image
        except ImportError:
            self.skipTest("需要 opencv 才能解碼驗證")
        self.login("t001")
        r = self.client.get("/dashboard/qr.png")
        self.assertEqual(r.mimetype, "image/png")
        img = np.array(Image.open(io.BytesIO(r.data)).convert("L"))
        val, _, _ = cv2.QRCodeDetector().detectAndDecode(img)
        self.assertEqual(val, "http://localhost/teacher/001")

    def test_public_qr_only_for_published(self):
        c = self.app.test_client()
        self.assertEqual(c.get("/teacher/001/qr.svg").status_code, 200)
        self.assertEqual(c.get("/teacher/003/qr.svg").status_code, 404)

    def test_qr_requires_login(self):
        self.assertEqual(self.app.test_client().get("/dashboard/qr.svg").status_code, 302)


class ImportTests(Base):
    def setUp(self):
        super().setUp()
        self.login("admin", "Admin1234")

    def upload(self, data, name):
        return self.post("/admin/import", {"file": (io.BytesIO(data), name)}, csrf_from="/admin/import")

    def payload(self, html):
        m = re.search(r'name="payload" value="([^"]*)"', html)
        import html as h
        return h.unescape(m.group(1))

    def test_csv_big5_preview_then_confirm(self):
        csv_text = ("代碼,姓名,帳號,Email,學校,縣市\n"
                    "010,王小明,,wang@ex.edu.tw,測試國小,台北市\n"
                    ",李小華,,,測試國中,\n"
                    "011,重複帳號,t001,,,\n"
                    "012,,t012,,,\n"
                    "013,壞信箱,t013,not-an-email,,\n")
        r = self.upload(csv_text.encode("cp950"), "名單.csv")
        html = r.get_data(as_text=True)
        self.assertIn("2 筆可建立", html)
        self.assertIn("帳號已存在", html)
        self.assertIn("缺少姓名", html)
        self.assertIn("Email 格式錯誤", html)
        self.assertEqual(self.db("SELECT COUNT(*) FROM users WHERE username='t010'")[0][0], 0)   # 預覽不寫入
        r = self.post("/admin/import/confirm", {"payload": self.payload(html)}, csrf_from="/admin/")
        html = r.get_data(as_text=True)
        self.assertEqual(len(re.findall(r'class="secret"', html)), 2)
        row = self.db("SELECT u.must_change_password, p.county, p.school_name FROM users u"
                      " JOIN teacher_profiles p ON p.user_id=u.id WHERE u.username='t010'")[0]
        self.assertEqual(tuple(row), (1, "臺北市", "測試國小"))
        auto = self.db("SELECT u.username FROM users u JOIN teacher_profiles p ON p.user_id=u.id"
                       " WHERE u.display_name='李小華'")[0][0]
        self.assertRegex(auto, r"^t\d{3}$")
        pw = re.findall(r'class="secret">([^<]+)<', html)[0]
        t = self.app.test_client()
        self.assertIn("/account/password", self.login("t010", pw, client=t).headers["Location"])

    def test_xlsx_import(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["teacher_code", "name", "username", "email", "school"])
        ws.append([20, "Excel 老師", "excel20", "x20@ex.edu.tw", "Excel 國小"])
        buf = io.BytesIO()
        wb.save(buf)
        html = self.upload(buf.getvalue(), "list.xlsx").get_data(as_text=True)
        self.assertIn("1 筆可建立", html)
        self.assertIn("020", html)          # Excel 數字 20 自動補成 020

    def test_confirm_revalidates_tampered_payload(self):
        payload = json.dumps([{"teacher_code": "001", "name": "冒用", "username": "t001"}])
        html = self.post("/admin/import/confirm", {"payload": payload}, csrf_from="/admin/").get_data(as_text=True)
        self.assertIn("帳號已存在", html)
        self.assertEqual(self.db("SELECT COUNT(*) FROM users WHERE display_name='冒用'")[0][0], 0)

    def test_teacher_cannot_import(self):
        t = self.app.test_client()
        self.login("t001", client=t)
        self.assertEqual(t.get("/admin/import").status_code, 403)

    def test_bad_file_type(self):
        r = self.upload(b"MZ\x90\x00", "virus.exe")
        self.assertEqual(r.status_code, 302)


class HistoryTests(Base):
    def test_restore_previous_version(self):
        self.login("t001")
        base = {"county": "新北市", "school_name": "新北市板橋區示範國民小學", "teacher_name": "林怡君"}
        self.post("/dashboard/profile", {**base, "slogan": "版本二"})
        self.post("/dashboard/publish")
        visitor = self.app.test_client()
        self.assertIn("版本二", visitor.get("/teacher/001").get_data(as_text=True))
        first = self.db("SELECT h.id FROM publish_history h JOIN teacher_profiles p ON p.id=h.profile_id"
                        " WHERE p.teacher_code='001' ORDER BY h.id")[0][0]
        self.post(f"/dashboard/history/{first}/restore")
        html = visitor.get("/teacher/001").get_data(as_text=True)
        self.assertNotIn("版本二", html)
        self.assertIn("讓每個孩子都成為會提問的學習者", html)
        # 草稿不受影響
        self.assertEqual(self.db("SELECT slogan FROM teacher_profiles WHERE teacher_code='001'")[0][0], "版本二")

    def test_cannot_restore_other_teachers_history(self):
        other = self.db("SELECT h.id FROM publish_history h JOIN teacher_profiles p ON p.id=h.profile_id"
                        " WHERE p.teacher_code='002'")[0][0]
        self.login("t001")
        self.assertEqual(self.post(f"/dashboard/history/{other}/restore").status_code, 404)


class ViewStatsTests(Base):
    def test_counts_once_per_browser_per_day_and_skips_owner_and_bots(self):
        def views():
            return self.db("SELECT COALESCE(SUM(views),0) FROM page_views v JOIN teacher_profiles p ON p.id=v.profile_id"
                           " WHERE p.teacher_code='002'")[0][0]
        start = views()
        ua = {"User-Agent": "Mozilla/5.0 (iPhone)"}
        v1 = self.app.test_client()
        v1.get("/teacher/002", headers=ua)
        v1.get("/teacher/002", headers=ua)
        self.assertEqual(views(), start + 1)
        self.app.test_client().get("/teacher/002", headers=ua)
        self.assertEqual(views(), start + 2)
        self.app.test_client().get("/teacher/002", headers={"User-Agent": "facebookexternalhit/1.1"})
        self.assertEqual(views(), start + 2)
        owner = self.app.test_client()
        self.login("t002", client=owner)
        owner.get("/teacher/002", headers=ua)
        self.assertEqual(views(), start + 2)
        self.app.test_client().get("/teacher/003", headers=ua)          # 草稿不計
        self.assertEqual(self.db("SELECT COUNT(*) FROM page_views v JOIN teacher_profiles p ON p.id=v.profile_id"
                                 " WHERE p.teacher_code='003'")[0][0], 0)

    def test_teacher_sees_only_own_stats(self):
        self.login("t001")
        html = self.client.get("/dashboard/").get_data(as_text=True)
        self.assertIn("公開頁瀏覽", html)
        self.assertNotIn("陳建宏", html)


class ValidityTests(Base):
    def test_expired_account_cannot_login_and_public_rule(self):
        self.db("UPDATE users SET valid_until='2000-01-01', keep_public_after_expiry=0 WHERE username='t002'")
        self.assertEqual(self.login("t002").status_code, 403)
        self.assertEqual(self.app.test_client().get("/teacher/002").status_code, 404)
        self.db("UPDATE users SET keep_public_after_expiry=1 WHERE username='t002'")
        self.assertEqual(self.app.test_client().get("/teacher/002").status_code, 200)

    def test_not_started_account(self):
        self.db("UPDATE users SET valid_from='2999-01-01' WHERE username='t003'")
        self.assertEqual(self.login("t003").status_code, 403)

    def test_expiry_kicks_existing_session(self):
        self.login("t001")
        self.assertEqual(self.client.get("/dashboard/").status_code, 200)
        self.db("UPDATE users SET valid_until='2000-01-01' WHERE username='t001'")
        self.assertEqual(self.client.get("/dashboard/").status_code, 302)

    def test_admin_sets_validity(self):
        self.login("admin", "Admin1234")
        uid = self.uid("t003")
        r = self.post(f"/admin/teachers/{uid}", {"username": "t003", "display_name": "張雅婷", "teacher_code": "003",
                                                 "valid_from": "2026-01-01", "valid_until": "2025-01-01"},
                      csrf_from="/admin/")
        self.assertEqual(r.status_code, 400)        # 結束早於開始
        self.post(f"/admin/teachers/{uid}", {"username": "t003", "display_name": "張雅婷", "teacher_code": "003",
                                             "valid_until": "2030-12-31"}, csrf_from="/admin/")
        self.assertEqual(self.db("SELECT valid_until FROM users WHERE id=%s", (uid,))[0][0], "2030-12-31")


class AnnouncementTests(Base):
    def test_announcement_window(self):
        self.login("admin", "Admin1234")
        self.post("/admin/announcements", {"title": "現在公告", "body": "內容"}, csrf_from="/admin/")
        self.post("/admin/announcements", {"title": "未來公告", "starts_on": "2999-01-01"}, csrf_from="/admin/")
        self.post("/admin/announcements", {"title": "過期公告", "ends_on": "2000-01-01"}, csrf_from="/admin/")
        self.post("/admin/announcements", {"action": "deadline", "deadline": "2099-12-31"}, csrf_from="/admin/")
        t = self.app.test_client()
        self.login("t001", client=t)
        html = t.get("/dashboard/").get_data(as_text=True)
        self.assertIn("現在公告", html)
        self.assertNotIn("未來公告", html)
        self.assertNotIn("過期公告", html)
        self.assertIn("2099-12-31", html)

    def test_teacher_cannot_post_announcement(self):
        self.login("t001")
        self.assertEqual(self.post("/admin/announcements", {"title": "x"}).status_code, 403)


class DashboardTests(Base):
    def test_filters_and_reminder_csv(self):
        self.login("admin", "Admin1234")
        html = self.client.get("/admin/?status=never").get_data(as_text=True)
        self.assertIn("黃志明", html)
        html = self.client.get("/admin/?status=complete").get_data(as_text=True)
        self.assertIn("林怡君", html)
        self.assertNotIn("黃志明", html)
        csv_text = self.client.get("/admin/reminders.csv?status=incomplete").get_data(as_text=True)
        self.assertIn("黃志明", csv_text)
        self.assertNotIn("林怡君", csv_text)

    def test_checklist_links_and_optional(self):
        self.login("t004")
        self.post("/account/password", {"current_password": "Demo1234", "new_password": "NewPass2026",
                                        "confirm_password": "NewPass2026"}, csrf_from="/account/password")
        html = self.client.get("/dashboard/").get_data(as_text=True)
        self.assertIn("前往填寫", html)
        self.assertIn("/dashboard/profile#avatar", html)
        self.assertIn("選填", html)

    def test_xlsx_export_sheets(self):
        from openpyxl import load_workbook
        self.login("admin", "Admin1234")
        wb = load_workbook(io.BytesIO(self.client.get("/admin/export.xlsx").data))
        self.assertEqual(wb.sheetnames, ["外牆與一頁網", "會場手冊", "課程內容", "學經歷與獲獎", "聯繫與進度（內部）", "瀏覽統計"])
        self.assertEqual(wb["會場手冊"]["D2"].value, "林怡君")


class BackupTests(Base):
    def test_backup_create_download_restore(self):
        from app.backup import restore_backup
        self.login("admin", "Admin1234")
        self.post("/admin/backups", csrf_from="/admin/")
        name = re.search(r"(backup-\d{8}-\d{6}\.zip)", self.client.get("/admin/backups").get_data(as_text=True)).group(1)
        r = self.client.get(f"/admin/backups/{name}")
        self.assertEqual(r.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(r.data)) as zf:
            self.assertIn("db/users.csv", zf.namelist())
            self.assertTrue(any(n.startswith("uploads/") for n in zf.namelist()))
        self.assertEqual(self.client.get("/admin/backups/../db/users.csv").status_code, 404)
        # 還原：先改資料，再還原回備份時的狀態
        self.db("UPDATE teacher_profiles SET slogan='備份後改的' WHERE teacher_code='001'")
        restore_backup(self.app.config, self.app.config["INSTANCE_DIR"] / "backups" / name)
        self.assertNotEqual(self.db("SELECT slogan FROM teacher_profiles WHERE teacher_code='001'")[0][0], "備份後改的")
        # 還原後自動編號要接在既有 id 之後，新增資料不可撞號
        self.db("INSERT INTO announcements (title, created_at, updated_at) VALUES ('還原後新增', '-', '-')")
        self.assertEqual(self.db("SELECT COUNT(*) FROM announcements")[0][0], 2)

    def test_backup_covers_every_table(self):
        from app.backup import TABLES
        tables = {r[0] for r in self.db("SELECT tablename FROM pg_tables WHERE schemaname = current_schema()")}
        self.assertEqual(set(TABLES), tables)

    def test_teacher_cannot_download_backup(self):
        self.login("t001")
        self.assertEqual(self.client.get("/admin/backups").status_code, 403)


class PublicPageTests(Base):
    def test_og_meta_and_no_private_data(self):
        html = self.client.get("/teacher/001").get_data(as_text=True)
        self.assertIn('property="og:title" content="林怡君｜新北市板橋區示範國民小學"', html)
        self.assertIn('property="og:image" content="http://localhost/files/', html)
        self.assertIn("data-lightbox", html)
        self.assertNotIn("0912-345-678", html)
        self.assertNotIn("yijun.lin@example.edu.tw", html)

    def test_thumbnails_follow_original_permissions(self):
        pub = self.db("SELECT f.thumb_name FROM uploaded_files f JOIN teacher_profiles p ON p.id=f.profile_id"
                      " WHERE p.teacher_code='001' AND f.category='class_photo'")[0][0]
        draft = self.db("SELECT f.thumb_name FROM uploaded_files f JOIN teacher_profiles p ON p.id=f.profile_id"
                        " WHERE p.teacher_code='003' AND f.thumb_name IS NOT NULL")[0][0]
        self.assertTrue(pub.startswith("t_") and pub.endswith(".webp"))
        c = self.app.test_client()
        r = c.get(f"/files/{pub}")
        self.assertEqual((r.status_code, r.mimetype), (200, "image/webp"))
        self.assertEqual(c.get(f"/files/{draft}").status_code, 404)

    def test_upload_creates_thumbnail(self):
        self.login("t003")
        self.post("/dashboard/photos", {"photos": (io.BytesIO(png_bytes((1600, 1200))), "big.png")})
        row = self.db("SELECT thumb_name FROM uploaded_files WHERE category='class_photo' ORDER BY id DESC LIMIT 1")[0]
        self.assertIsNotNone(row[0])
        from PIL import Image
        with Image.open(self.app.config["UPLOAD_FOLDER"] / row[0]) as im:
            self.assertLessEqual(max(im.size), 720)

    def test_home_highlights_and_sitemap(self):
        html = self.client.get("/").get_data(as_text=True)
        self.assertIn("教師列表", html)
        sm = self.client.get("/sitemap.xml").get_data(as_text=True)
        self.assertIn("/teacher/001", sm)
        self.assertNotIn("/teacher/003", sm)
