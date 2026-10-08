"""管理員後台：帳號管理、批次匯入、進度追蹤、提醒、公告、匯出、備份。全部路由都需 admin 角色。"""
import csv
import io
import json
import re
from datetime import datetime, timedelta

import psycopg
from flask import (Blueprint, Response, abort, current_app, flash, g, redirect, render_template, request,
                   send_file, url_for)

from .backup import NAME_RE as BACKUP_NAME_RE
from .backup import create_backup, list_backups
from .db import execute, get_db, get_setting, insert, local_today, parse_ts, query, set_setting, ts
from .mail import send_mail
from .profiles import checklist, completeness, has_unpublished_changes, load_courses, load_items
from .security import (admin_required, audit, destroy_user_sessions, generate_temp_password,
                       hash_password, password_problems)
from .services import public_url, site_root, view_stats

bp = Blueprint("admin", __name__)

USERNAME_RE = re.compile(r"^[A-Za-z0-9_.-]{3,32}$")
CODE_RE = re.compile(r"^[A-Za-z0-9_-]{1,20}$")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


@bp.before_request
@admin_required
def _guard():
    """整個 /admin 藍圖統一檢查：未登入導向登入頁，教師角色一律 403。"""
    return None


def teacher_or_404(uid):
    row = query("SELECT u.*, p.id AS profile_id, p.teacher_code, p.status, p.updated_at AS profile_updated_at,"
                " p.published_at, p.school_name, p.teacher_name, p.county"
                " FROM users u JOIN teacher_profiles p ON p.user_id = u.id"
                " WHERE u.id = %s AND u.role = 'teacher'", (uid,), one=True)
    if row is None:
        abort(404)
    return row


def next_teacher_code(taken=()):
    nums = [int(r["teacher_code"]) for r in query("SELECT teacher_code FROM teacher_profiles")
            if r["teacher_code"].isdigit()]
    nums += [int(c) for c in taken if c.isdigit()]
    return f"{(max(nums) + 1) if nums else 1:03d}"


# ================================================================ 進度管理 Dashboard
FILTERS = [
    ("", "全部"), ("complete", "已完成"), ("incomplete", "未完成"), ("low", "完成度 < 50%"),
    ("published", "已發布"), ("unpublished", "未發布"), ("changes", "有未發布修改"),
    ("must_change", "待改密碼"), ("never", "未登入過"), ("stale", "7 天以上未更新"), ("inactive", "已停用"),
]


def teacher_overview_rows():
    """所有老師的進度資料（Dashboard、提醒名單、匯出共用）。"""
    stale_before = (parse_ts(ts()) - timedelta(days=7)).strftime("%Y-%m-%d %H:%M:%S")
    views = view_stats()
    rows = query("SELECT u.id, u.username, u.display_name, u.email, u.is_active, u.last_login_at,"
                 " u.must_change_password, u.valid_from, u.valid_until,"
                 " p.id AS profile_id, p.teacher_code, p.school_name, p.county, p.teacher_name, p.status,"
                 " p.updated_at, p.published_at, p.contact_email"
                 " FROM users u JOIN teacher_profiles p ON p.user_id = u.id WHERE u.role = 'teacher'"
                 " ORDER BY p.teacher_code")
    out = []
    for r in rows:
        items = checklist(r["profile_id"])
        req = [i for i in items if i["required"]]
        pct = round(sum(i["done"] for i in req) * 100 / len(req))
        out.append(dict(r, pct=pct, missing=[i["label"] for i in req if not i["done"]],
                        has_changes=has_unpublished_changes(r), stale=r["updated_at"] < stale_before,
                        views=views.get(r["profile_id"], {"total": 0, "d7": 0, "d30": 0})))
    return out


def apply_filter(rows, key, q=""):
    tests = {
        "complete": lambda r: r["pct"] == 100,
        "incomplete": lambda r: r["pct"] < 100,
        "low": lambda r: r["pct"] < 50,
        "published": lambda r: r["status"] == "published",
        "unpublished": lambda r: r["status"] != "published",
        "changes": lambda r: r["has_changes"],
        "must_change": lambda r: r["must_change_password"],
        "never": lambda r: not r["last_login_at"],
        "stale": lambda r: r["stale"],
        "inactive": lambda r: not r["is_active"],
    }
    if key in tests:
        rows = [r for r in rows if tests[key](r)]
    if key != "inactive" and key:
        rows = [r for r in rows if r["is_active"]]
    if q:
        ql = q.lower()
        rows = [r for r in rows if ql in " ".join(str(r[k] or "") for k in (
            "username", "display_name", "teacher_name", "school_name", "teacher_code", "county")).lower()]
    return rows


@bp.route("/")
def dashboard():
    q = (request.args.get("q") or "").strip()[:50]
    status = request.args.get("status", "")
    if status not in dict(FILTERS):
        status = ""
    all_rows = teacher_overview_rows()
    active = [r for r in all_rows if r["is_active"]]
    stats = {
        "total": len(all_rows),
        "published": sum(1 for r in active if r["status"] == "published"),
        "unpublished": sum(1 for r in active if r["status"] != "published"),
        "low": sum(1 for r in active if r["pct"] < 50),
        "never": sum(1 for r in active if not r["last_login_at"]),
        "must_change": sum(1 for r in active if r["must_change_password"]),
        "stale": sum(1 for r in active if r["stale"]),
        "inactive": sum(1 for r in all_rows if not r["is_active"]),
        "avg": round(sum(r["pct"] for r in active) / len(active)) if active else 0,
    }
    counts = {k: len(apply_filter(all_rows, k)) for k, _ in FILTERS}
    teachers = apply_filter(all_rows, status, q)
    return render_template("admin/dashboard.html", teachers=teachers, stats=stats, filters=FILTERS,
                           counts=counts, current={"q": q, "status": status},
                           deadline=get_setting("deadline"), last_backup=get_setting("last_backup_at"),
                           smtp=bool(current_app.config["SMTP_HOST"]))


# ================================================================ 提醒名單
@bp.route("/reminders.csv")
def reminders_csv():
    rows = apply_filter(teacher_overview_rows(), request.args.get("status", "incomplete"),
                        (request.args.get("q") or "").strip())
    out = io.StringIO()
    out.write("﻿")
    w = csv.writer(out)
    w.writerow(["網址代碼", "姓名", "學校", "帳號", "Email", "完成度", "尚缺項目", "發布狀態", "最後更新", "最後登入"])
    for r in rows:
        w.writerow([_csv_safe(v) for v in [
            r["teacher_code"], r["teacher_name"] or r["display_name"], r["school_name"], r["username"],
            r["email"] or r["contact_email"], f"{r['pct']}%", "、".join(r["missing"]),
            "已發布" if r["status"] == "published" else "未發布",
            _local(r["updated_at"]), _local(r["last_login_at"]) or "未登入"]])
    audit("export_reminders")
    return Response(out.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": "attachment; filename=reminders.csv",
                             "Cache-Control": "no-store"})


@bp.route("/reminders/send", methods=["POST"])
def reminders_send():
    if not current_app.config["SMTP_HOST"]:
        flash("尚未設定寄信（SMTP），無法寄送提醒。請先在 env.json 設定 SMTP_HOST 等參數，或匯出名單自行聯繫。", "error")
        return redirect(url_for("admin.dashboard"))
    status = request.form.get("status", "incomplete")
    rows = apply_filter(teacher_overview_rows(), status, (request.form.get("q") or "").strip())
    deadline = get_setting("deadline")
    login = site_root() + url_for("auth.login")
    sent = skipped = 0
    for r in rows:
        to = r["email"] or r["contact_email"]
        if not to:
            skipped += 1
            continue
        body = (f"{r['teacher_name'] or r['display_name']} 老師您好：\n\n"
                f"提醒您「{current_app.config['SITE_NAME']}」的資料目前完成度為 {r['pct']}%"
                + (f"，尚缺：{'、'.join(r['missing'])}" if r["missing"] else "")
                + ("，目前尚未發布" if r["status"] != "published" else "") + "。\n"
                + (f"填寫截止日：{deadline}\n" if deadline else "")
                + f"\n登入網址：{login}\n帳號：{r['username']}\n\n謝謝您的協助！")
        if send_mail(to, f"[{current_app.config['SITE_NAME']}] 資料填寫提醒", body):
            sent += 1
        else:
            skipped += 1
    audit("reminders_sent", detail=f"sent={sent} skipped={skipped} filter={status}")
    flash(f"已寄出 {sent} 封提醒信" + (f"，{skipped} 位因無 Email 或寄送失敗未寄出" if skipped else "") + "。", "success")
    return redirect(url_for("admin.dashboard", status=status))


# ================================================================ 單筆帳號
def _read_account_form(existing=None):
    data = {
        "username": (request.form.get("username") or "").strip(),
        "display_name": (request.form.get("display_name") or "").strip()[:50],
        "email": (request.form.get("email") or "").strip()[:120] or None,
        "teacher_code": (request.form.get("teacher_code") or "").strip(),
        "valid_from": (request.form.get("valid_from") or "").strip() or None,
        "valid_until": (request.form.get("valid_until") or "").strip() or None,
        "keep_public_after_expiry": 1 if request.form.get("keep_public_after_expiry") else 0,
    }
    errors = []
    if not USERNAME_RE.match(data["username"]):
        errors.append("帳號需為 3–32 字元的英數字（可含 _ . -）")
    if not data["display_name"]:
        errors.append("請填寫教師姓名")
    if data["email"] and not EMAIL_RE.match(data["email"]):
        errors.append("Email 格式不正確")
    if not CODE_RE.match(data["teacher_code"]):
        errors.append("網址代碼需為 1–20 字元英數字，例如 001")
    for k, label in (("valid_from", "開始日期"), ("valid_until", "結束日期")):
        if data[k] and not DATE_RE.match(data[k]):
            errors.append(f"{label}格式需為 YYYY-MM-DD")
    if data["valid_from"] and data["valid_until"] and data["valid_from"] > data["valid_until"]:
        errors.append("帳號結束日期不可早於開始日期")
    uid = existing["id"] if existing else -1
    if query("SELECT 1 FROM users WHERE lower(username) = lower(%s) AND id != %s", (data["username"], uid), one=True):
        errors.append("此帳號已被使用")
    if query("SELECT 1 FROM teacher_profiles WHERE teacher_code = %s AND user_id != %s",
             (data["teacher_code"], uid), one=True):
        errors.append("此網址代碼已被使用")
    if data["email"] and query("SELECT 1 FROM users WHERE lower(email) = lower(%s) AND id != %s",
                               (data["email"], uid), one=True):
        errors.append("此 Email 已被其他帳號使用")
    return data, errors


def _create_teacher(data, password):
    uid = insert(
        "INSERT INTO users (username, password_hash, role, display_name, email, is_active, must_change_password,"
        " valid_from, valid_until, keep_public_after_expiry, created_at, updated_at)"
        " VALUES (%s, %s, 'teacher', %s, %s, 1, 1, %s, %s, %s, %s, %s)",
        (data["username"], hash_password(password), data["display_name"], data["email"],
         data.get("valid_from"), data.get("valid_until"), data.get("keep_public_after_expiry", 1), ts(), ts()))
    execute("INSERT INTO teacher_profiles (user_id, teacher_code, teacher_name, school_name, county,"
            " contact_email, created_at, updated_at) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)",
            (uid, data["teacher_code"], data["display_name"], data.get("school"), data.get("county"),
             data["email"], ts(), ts()))
    return uid


@bp.route("/teachers/new", methods=["GET", "POST"])
def teacher_new():
    if request.method == "POST":
        data, errors = _read_account_form()
        password = request.form.get("password") or ""
        if password:
            errors += [f"初始密碼{p}" for p in password_problems(password, data["username"])]
        if errors:
            for e in errors:
                flash(e, "error")
            return render_template("admin/teacher_form.html", form=data, teacher=None), 400
        temp = password or generate_temp_password()
        db = get_db()
        try:
            with db.transaction():   # 帳號與教師資料在同一個交易內建立
                uid = _create_teacher(data, temp)
        except psycopg.IntegrityError:
            flash("帳號或網址代碼重複，請重新確認。", "error")
            return render_template("admin/teacher_form.html", form=data, teacher=None), 400
        audit("teacher_created", f"user:{uid}", data["username"])
        t = teacher_or_404(uid)
        return render_template("admin/credentials.html", created=True, rows=[_cred_row(t, temp)])
    return render_template("admin/teacher_form.html",
                           form={"teacher_code": next_teacher_code(), "keep_public_after_expiry": 1}, teacher=None)


def _cred_row(t, password):
    return {"name": t["teacher_name"] or t["display_name"], "school": t["school_name"] or "",
            "username": t["username"], "password": password, "code": t["teacher_code"],
            "email": t["email"] or "", "public_url": public_url(t["teacher_code"])}


def _teacher_extras(teacher):
    items = checklist(teacher["profile_id"])
    req = [i for i in items if i["required"]]
    return {"checklist": items, "pct": round(sum(i["done"] for i in req) * 100 / len(req)),
            "views": view_stats(teacher["profile_id"])}


@bp.route("/teachers/<int:uid>", methods=["GET", "POST"])
def teacher_edit(uid):
    teacher = teacher_or_404(uid)
    if request.method == "POST":
        data, errors = _read_account_form(existing=teacher)
        if errors:
            for e in errors:
                flash(e, "error")
            return render_template("admin/teacher_form.html", form=data, teacher=teacher,
                                   **_teacher_extras(teacher)), 400
        execute("UPDATE users SET username = %s, display_name = %s, email = %s, valid_from = %s, valid_until = %s,"
                " keep_public_after_expiry = %s, updated_at = %s WHERE id = %s",
                (data["username"], data["display_name"], data["email"], data["valid_from"], data["valid_until"],
                 data["keep_public_after_expiry"], ts(), uid))
        execute("UPDATE teacher_profiles SET teacher_code = %s, updated_at = %s WHERE user_id = %s",
                (data["teacher_code"], ts(), uid))
        audit("teacher_updated", f"user:{uid}")
        flash("教師帳號已更新。", "success")
        return redirect(url_for("admin.dashboard"))
    return render_template("admin/teacher_form.html", form=dict(teacher), teacher=teacher,
                           **_teacher_extras(teacher))


@bp.route("/teachers/<int:uid>/toggle-active", methods=["POST"])
def teacher_toggle(uid):
    teacher = teacher_or_404(uid)
    new = 0 if teacher["is_active"] else 1
    execute("UPDATE users SET is_active = %s, updated_at = %s WHERE id = %s", (new, ts(), uid))
    if not new:
        destroy_user_sessions(uid)          # 停用立即生效
    audit("teacher_enabled" if new else "teacher_disabled", f"user:{uid}")
    flash(f"已{'啟用' if new else '停用'}「{teacher['display_name']}」。"
          + ("" if new else "停用期間無法登入，公開頁面也會暫時隱藏。"), "success")
    return redirect(request.form.get("back") == "edit" and url_for("admin.teacher_edit", uid=uid)
                    or url_for("admin.dashboard"))


@bp.route("/teachers/<int:uid>/reset-password", methods=["POST"])
def teacher_reset_password(uid):
    teacher = teacher_or_404(uid)
    temp = generate_temp_password()
    execute("UPDATE users SET password_hash = %s, must_change_password = 1, failed_login_count = 0,"
            " locked_until = NULL, updated_at = %s WHERE id = %s", (hash_password(temp), ts(), uid))
    destroy_user_sessions(uid)
    audit("teacher_password_reset", f"user:{uid}")
    # 直接回應頁面（不 redirect），臨時密碼只顯示這一次、不寫入任何紀錄
    return render_template("admin/credentials.html", created=False, rows=[_cred_row(teacher, temp)])


# ================================================================ 批次匯入
FIELD_ALIASES = {
    "teacher_code": ("teacher_code", "code", "代碼", "網址代碼", "編號"),
    "name": ("name", "姓名", "教師姓名", "講師姓名"),
    "username": ("username", "帳號", "登入帳號"),
    "email": ("email", "e-mail", "電子郵件", "信箱", "電子信箱"),
    "school": ("school", "學校", "學校全銜", "學校名稱"),
    "county": ("county", "縣市"),
}
MAX_IMPORT_ROWS = 200


def _parse_table(file_storage):
    """讀取 xlsx 或 csv，回傳 [[cell,...], ...]。CSV 自動判斷 UTF-8／Big5（Excel 另存 CSV 常見）。"""
    name = (file_storage.filename or "").lower()
    data = file_storage.read(2 * 1024 * 1024 + 1)
    if len(data) > 2 * 1024 * 1024:
        raise ValueError("檔案超過 2 MB")
    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("伺服器未安裝 openpyxl，請改用 CSV，或執行 pip install openpyxl") from exc
        if not data.startswith(b"PK"):
            raise ValueError("不是有效的 Excel 檔案")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [["" if c is None else str(c).strip() for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    if name.endswith(".csv"):
        for enc in ("utf-8-sig", "cp950", "big5hkscs"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("無法辨識 CSV 的文字編碼，請另存為「CSV UTF-8」")
        return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]
    raise ValueError("僅接受 .xlsx 或 .csv 檔案")


def _rows_from_table(table):
    table = [r for r in table if any(c for c in r)]
    if not table:
        raise ValueError("檔案是空的")
    header = [h.strip().lower() for h in table[0]]
    index = {}
    for field, aliases in FIELD_ALIASES.items():
        for i, h in enumerate(header):
            if h in aliases:
                index[field] = i
                break
    if "name" not in index:
        raise ValueError("找不到「姓名」(name) 欄位，請使用範本的標題列")
    rows = []
    for raw in table[1:MAX_IMPORT_ROWS + 1]:
        get = lambda f: (raw[index[f]] if f in index and index[f] < len(raw) else "").strip()  # noqa: E731
        code = get("teacher_code")
        if code.isdigit():
            code = code.zfill(3)            # Excel 會把 001 變成 1
        rows.append({"teacher_code": code, "name": get("name")[:50], "username": get("username"),
                     "email": get("email")[:120], "school": get("school")[:100], "county": get("county")[:10]})
    if len(table) - 1 > MAX_IMPORT_ROWS:
        raise ValueError(f"一次最多匯入 {MAX_IMPORT_ROWS} 筆")
    return rows


def _validate_import(rows):
    """補齊空白欄位並逐筆檢查。回傳每筆的 errors；不會寫入資料庫。"""
    from .profiles import COUNTIES
    existing_users = {r["username"].lower() for r in query("SELECT username FROM users")}
    existing_codes = {r["teacher_code"] for r in query("SELECT teacher_code FROM teacher_profiles")}
    existing_emails = {(r["email"] or "").lower() for r in query("SELECT email FROM users") if r["email"]}
    seen_u, seen_c, seen_e, auto_codes = set(), set(), set(), []
    for r in rows:
        if not r["teacher_code"]:
            r["teacher_code"] = next_teacher_code(taken=list(seen_c) + auto_codes)
            auto_codes.append(r["teacher_code"])
            r["auto_code"] = True
        if not r["username"] and CODE_RE.match(r["teacher_code"]):
            r["username"] = "t" + r["teacher_code"]
            r["auto_username"] = True
        errs = []
        if not r["name"]:
            errs.append("缺少姓名")
        if not USERNAME_RE.match(r["username"]):
            errs.append("帳號需為 3–32 個英數字")
        elif r["username"].lower() in existing_users:
            errs.append("帳號已存在")
        elif r["username"].lower() in seen_u:
            errs.append("帳號在檔案中重複")
        if not CODE_RE.match(r["teacher_code"]):
            errs.append("代碼格式錯誤")
        elif r["teacher_code"] in existing_codes:
            errs.append("代碼已存在")
        elif r["teacher_code"] in seen_c:
            errs.append("代碼在檔案中重複")
        if r["email"]:
            if not EMAIL_RE.match(r["email"]):
                errs.append("Email 格式錯誤")
            elif r["email"].lower() in existing_emails:
                errs.append("Email 已被使用")
            elif r["email"].lower() in seen_e:
                errs.append("Email 在檔案中重複")
        if r["county"] and r["county"].replace("台", "臺") in COUNTIES:
            r["county"] = r["county"].replace("台", "臺")
        elif r["county"]:
            errs.append("縣市名稱無法辨識")
        r["errors"] = errs
        seen_u.add(r["username"].lower())
        seen_c.add(r["teacher_code"])
        if r["email"]:
            seen_e.add(r["email"].lower())
    return rows


@bp.route("/import", methods=["GET", "POST"])
def import_teachers():
    if request.method == "POST":
        f = request.files.get("file")
        if not f or not f.filename:
            flash("請選擇要匯入的檔案。", "error")
            return redirect(url_for("admin.import_teachers"))
        try:
            rows = _validate_import(_rows_from_table(_parse_table(f)))
        except ValueError as e:
            flash(f"無法讀取檔案：{e}", "error")
            return redirect(url_for("admin.import_teachers"))
        except Exception:     # noqa: BLE001 — 損毀的 Excel 等
            current_app.logger.exception("import parse failed")
            flash("無法讀取檔案，請確認是範本格式的 .xlsx 或 .csv。", "error")
            return redirect(url_for("admin.import_teachers"))
        ok = [r for r in rows if not r["errors"]]
        payload = json.dumps([{k: r[k] for k in ("teacher_code", "name", "username", "email", "school", "county")}
                              for r in ok], ensure_ascii=False)
        return render_template("admin/import_preview.html", rows=rows, ok_count=len(ok),
                               bad_count=len(rows) - len(ok), payload=payload)
    return render_template("admin/import.html")


@bp.route("/import/confirm", methods=["POST"])
def import_confirm():
    try:
        rows = json.loads(request.form.get("payload") or "[]")
        assert isinstance(rows, list) and len(rows) <= MAX_IMPORT_ROWS
        rows = [{k: str(r.get(k) or "").strip() for k in ("teacher_code", "name", "username", "email", "school",
                                                         "county")} for r in rows]
    except (ValueError, AssertionError, AttributeError):
        abort(400)
    rows = _validate_import(rows)          # 送出時再完整檢查一次（不信任前端資料）
    created, failed = [], []
    db = get_db()
    for r in rows:
        if r["errors"]:
            failed.append(r)
            continue
        temp = generate_temp_password()
        data = {"username": r["username"], "display_name": r["name"], "email": r["email"] or None,
                "teacher_code": r["teacher_code"], "school": r["school"] or None, "county": r["county"] or None}
        try:
            with db.transaction():
                uid = _create_teacher(data, temp)
        except psycopg.IntegrityError:
            r["errors"] = ["帳號或代碼已被使用"]
            failed.append(r)
            continue
        created.append(_cred_row(teacher_or_404(uid), temp))
    audit("teachers_imported", detail=f"created={len(created)} failed={len(failed)}")
    return render_template("admin/credentials.html", created=True, rows=created, failed=failed, batch=True)


@bp.route("/import/template.csv")
def import_template():
    out = io.StringIO()
    out.write("﻿")
    w = csv.writer(out)
    w.writerow(["teacher_code", "name", "username", "email", "school", "county"])
    w.writerow(["005", "王小明", "t005", "wang@example.edu.tw", "新北市○○區○○國民小學", "新北市"])
    w.writerow(["", "李小華", "", "", "臺中市○○國民中學", "臺中市"])
    return Response(out.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": "attachment; filename=teacher_import_template.csv"})


# ================================================================ 公告與截止日
@bp.route("/announcements", methods=["GET", "POST"])
def announcements():
    if request.method == "POST":
        action = request.form.get("action")
        if action == "deadline":
            d = (request.form.get("deadline") or "").strip()
            if d and not DATE_RE.match(d):
                flash("截止日格式需為 YYYY-MM-DD", "error")
            else:
                set_setting("deadline", d or None)
                audit("deadline_set", detail=d)
                flash("已更新填寫截止日。" if d else "已清除填寫截止日。", "success")
            return redirect(url_for("admin.announcements"))
        aid = request.form.get("id", type=int)
        if action == "delete" and aid:
            execute("DELETE FROM announcements WHERE id = %s", (aid,))
            audit("announcement_deleted", f"announcement:{aid}")
            flash("公告已刪除。", "success")
            return redirect(url_for("admin.announcements"))
        title = (request.form.get("title") or "").strip()[:100]
        body = (request.form.get("body") or "").strip()[:2000]
        starts = (request.form.get("starts_on") or "").strip() or None
        ends = (request.form.get("ends_on") or "").strip() or None
        pinned = 1 if request.form.get("is_pinned") else 0
        errors = []
        if not title:
            errors.append("請填寫公告標題")
        for v in (starts, ends):
            if v and not DATE_RE.match(v):
                errors.append("日期格式需為 YYYY-MM-DD")
        if starts and ends and starts > ends:
            errors.append("結束日期不可早於開始日期")
        if errors:
            for e in errors:
                flash(e, "error")
        elif aid:
            execute("UPDATE announcements SET title = %s, body = %s, starts_on = %s, ends_on = %s, is_pinned = %s,"
                    " updated_at = %s WHERE id = %s", (title, body, starts, ends, pinned, ts(), aid))
            audit("announcement_updated", f"announcement:{aid}")
            flash("公告已更新。", "success")
        else:
            aid = insert("INSERT INTO announcements (title, body, starts_on, ends_on, is_pinned, created_by,"
                         " created_at, updated_at) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)",
                         (title, body, starts, ends, pinned, g.user["id"], ts(), ts()))
            audit("announcement_created", f"announcement:{aid}")
            flash("公告已發布，老師登入後會在總覽頁看到。", "success")
        return redirect(url_for("admin.announcements"))
    edit = None
    if request.args.get("edit", type=int):
        edit = query("SELECT * FROM announcements WHERE id = %s", (request.args.get("edit", type=int),), one=True)
    rows = query("SELECT * FROM announcements ORDER BY is_pinned DESC, id DESC")
    return render_template("admin/announcements.html", rows=rows, edit=edit, today=local_today(),
                           deadline=get_setting("deadline"))


# ================================================================ 匯出
def _csv_safe(v):
    """防 CSV/Excel 公式注入：以 = + - @ 開頭的字串前面加單引號。"""
    if v is None:
        return ""
    s = str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@", "\t", "\r") else s


def _local(value):
    dt = parse_ts(value)
    return (dt + timedelta(hours=8)).strftime("%Y-%m-%d %H:%M") if dt else ""


@bp.route("/export.csv")
def export_csv():
    """匯出會場手冊 / 外牆所需欄位（Excel 可直接開啟的 UTF-8 BOM CSV）。內含聯絡資料，僅管理員可下載。"""
    out = io.StringIO()
    out.write("﻿")
    w = csv.writer(out)
    w.writerow(["網址代碼", "縣市", "學校全銜", "講師姓名", "職稱", "手機號碼", "電子郵件", "Slogan",
                "授課年級", "學生數", "科目/領域", "課程名稱", "單元", "發布狀態", "完成度", "最後修改"])
    for p in query("SELECT p.* FROM teacher_profiles p JOIN users u ON u.id = p.user_id"
                   " WHERE u.role = 'teacher' ORDER BY p.teacher_code"):
        pct, _ = completeness(p["id"])
        courses = load_courses(p["id"]) or [{}]
        for c in courses:
            w.writerow([_csv_safe(v) for v in [
                p["teacher_code"], p["county"], p["school_name"], p["teacher_name"], p["job_title"],
                p["phone"], p["contact_email"], p["slogan"], c.get("grade"), c.get("student_count"),
                c.get("subject"), c.get("title"), c.get("unit"),
                "已發布" if p["status"] == "published" else "草稿", f"{pct}%", _local(p["updated_at"])]])
    audit("export_csv")
    return Response(out.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": "attachment; filename=teachers_export.csv",
                             "Cache-Control": "no-store"})


@bp.route("/export.xlsx")
def export_xlsx():
    """完整 Excel：依『運用處』分工作表（外牆／一頁網、會場手冊、課程內容、學經歷與獲獎、填寫進度、瀏覽統計）。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        flash("伺服器未安裝 openpyxl，請執行 pip install -r requirements.txt 後再試，或先使用 CSV 匯出。", "error")
        return redirect(url_for("admin.dashboard"))
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F5F8B")

    def sheet(title, headers, rows, widths=None, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append([_csv_safe(v) if isinstance(v, str) else v for v in r])
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(40, len(h) * 2 + 4)))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        return ws

    profiles = query("SELECT p.*, u.username, u.is_active, u.last_login_at FROM teacher_profiles p"
                     " JOIN users u ON u.id = p.user_id WHERE u.role = 'teacher' ORDER BY p.teacher_code")
    overview = {r["profile_id"]: r for r in teacher_overview_rows()}
    sheet("外牆與一頁網", ["代碼", "縣市", "學校全銜", "講師姓名", "職稱", "Slogan", "教育理念", "教學專長",
                       "公開網址", "發布狀態"],
          [[p["teacher_code"], p["county"], p["school_name"], p["teacher_name"], p["job_title"], p["slogan"],
            p["philosophy"], "、".join(s["name"] for s in load_items("specialties", p["id"])),
            public_url(p["teacher_code"]), "已發布" if p["status"] == "published" else "未發布"] for p in profiles],
          {"教育理念": 50, "公開網址": 40, "學校全銜": 30}, first=True)
    hand, course_rows = [], []
    for p in profiles:
        for c in load_courses(p["id"]):
            hand.append([p["teacher_code"], p["county"], p["school_name"], p["teacher_name"], c["grade"],
                         c["student_count"], c["subject"], c["title"]])
            course_rows.append([p["teacher_code"], p["teacher_name"], c["title"], c["unit"], c["objectives"],
                                c["intro"], c["four_learning"], c["ai_usage"],
                                "、".join(lp["title"] for lp in c["plans"])])
    sheet("會場手冊", ["代碼", "縣市", "學校全銜", "講師姓名", "授課年級", "學生數", "科目/領域", "課程名稱"], hand,
          {"學校全銜": 30, "課程名稱": 36})
    sheet("課程內容", ["代碼", "講師姓名", "課程名稱", "單元", "學習目標", "課程介紹", "四學運用", "AI 運用", "教案檔案"],
          course_rows, {k: 45 for k in ("學習目標", "課程介紹", "四學運用", "AI 運用")})
    bg = []
    for p in profiles:
        for e in load_items("educations", p["id"]):
            bg.append([p["teacher_code"], p["teacher_name"], "學歷", f"{e['school']} {e['department'] or ''}",
                       e["degree"], e["period"]])
        for e in load_items("experiences", p["id"]):
            bg.append([p["teacher_code"], p["teacher_name"], "經歷", e["organization"], e["role"], e["period"]])
        for a in load_items("awards", p["id"]):
            bg.append([p["teacher_code"], p["teacher_name"], "獲獎", a["title"], a["issuer"], a["award_date"]])
        for c in load_items("certifications", p["id"]):
            bg.append([p["teacher_code"], p["teacher_name"], "認證", c["name"], c["issuer"], c["cert_date"]])
    sheet("學經歷與獲獎", ["代碼", "講師姓名", "類別", "名稱", "職務/學位/單位", "期間/日期"], bg, {"名稱": 40})
    sheet("聯繫與進度（內部）", ["代碼", "講師姓名", "帳號", "手機", "Email", "完成度", "尚缺項目", "發布狀態",
                         "最後更新", "最後登入", "帳號狀態"],
          [[p["teacher_code"], p["teacher_name"], p["username"], p["phone"], p["contact_email"],
            f"{overview[p['id']]['pct']}%", "、".join(overview[p["id"]]["missing"]),
            "已發布" if p["status"] == "published" else "未發布", _local(p["updated_at"]),
            _local(p["last_login_at"]) or "未登入", "啟用" if p["is_active"] else "停用"] for p in profiles],
          {"尚缺項目": 40})
    sheet("瀏覽統計", ["代碼", "講師姓名", "近 7 天", "近 30 天", "總瀏覽"],
          [[p["teacher_code"], p["teacher_name"], overview[p["id"]]["views"]["d7"],
            overview[p["id"]]["views"]["d30"], overview[p["id"]]["views"]["total"]] for p in profiles])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    audit("export_xlsx")
    stamp = (datetime.utcnow() + timedelta(hours=8)).strftime("%Y%m%d")
    resp = send_file(buf, as_attachment=True, download_name=f"teachers_{stamp}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp.headers["Cache-Control"] = "no-store"
    return resp


# ================================================================ 備份
@bp.route("/backups", methods=["GET", "POST"])
def backups():
    if request.method == "POST":
        try:
            path = create_backup(current_app.config)
        except Exception:     # noqa: BLE001
            current_app.logger.exception("backup failed")
            flash("備份失敗，請查看伺服器紀錄。", "error")
            return redirect(url_for("admin.backups"))
        audit("backup_created", detail=path.name)
        flash(f"已建立備份 {path.name}。建議下載一份存放到其他地方（例如雲端硬碟）。", "success")
        return redirect(url_for("admin.backups"))
    return render_template("admin/backups.html", rows=list_backups(current_app.config),
                           last_backup=get_setting("last_backup_at"))


@bp.route("/backups/<name>")
def backup_download(name):
    if not BACKUP_NAME_RE.match(name):
        abort(404)
    path = current_app.config["INSTANCE_DIR"] / "backups" / name
    if not path.is_file():
        abort(404)
    audit("backup_downloaded", detail=name)
    resp = send_file(path, as_attachment=True, download_name=name, mimetype="application/zip")
    resp.headers["Cache-Control"] = "no-store"
    return resp


# ================================================================ 操作紀錄
@bp.route("/audit")
def audit_log():
    rows = query("SELECT a.*, u.username FROM audit_logs a LEFT JOIN users u ON u.id = a.actor_id"
                 " ORDER BY a.id DESC LIMIT 300")
    return render_template("admin/audit.html", rows=rows)
